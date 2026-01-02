import streamlit as st
import google.generativeai as genai
from pypdf import PdfReader
from docx import Document
from io import BytesIO
from duckduckgo_search import DDGS
import requests
from bs4 import BeautifulSoup
import time

# Tenta importar a biblioteca legislativa local
try:
    import legislacao
except ImportError:
    st.error("⚠️ Ficheiro 'legislacao.py' não encontrado. Cria-o na mesma pasta do app.py.")
    st.stop()

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(
    page_title="Análise Ambiental IA",
    page_icon="🌿",
    layout="wide"
)

# --- GESTÃO DE ESTADO (SESSION STATE) ---
if 'uploader_key' not in st.session_state:
    st.session_state.uploader_key = 0

def limpar_dados():
    """Função para limpar os ficheiros e reiniciar a app"""
    st.session_state.uploader_key += 1
    st.rerun()

# --- ESTILO CSS ---
st.markdown("""
<style>
    .stButton>button {
        width: 100%; 
        border-radius: 8px; 
        height: 3em;
        font-weight: bold;
    }
    .reportview-container { margin-top: -2em; }
    h1 { color: #2e7d32; }
    .stExpander { border: 1px solid #ddd; border-radius: 5px; }
</style>
""", unsafe_allow_html=True)

# --- CABEÇALHO ---
col1, col2 = st.columns([1, 6])
with col1:
    st.markdown("# 🌿")
with col2:
    st.title("Análise Ambiental")
    st.caption("Protocolo PATE v4.2 | Compliance, EIA e Sustentabilidade | Seletor de Modelos")

# --- FUNÇÃO PARA LISTAR MODELOS ---
def get_available_models(api_key):
    """Obtém lista de modelos disponíveis na API Key do utilizador."""
    try:
        genai.configure(api_key=api_key)
        models = []
        for m in genai.list_models():
            # Filtra apenas modelos que geram texto (chat)
            if 'generateContent' in m.supported_generation_methods:
                models.append(m.name)
        return models
    except Exception as e:
        return []

# --- SIDEBAR: CONFIGURAÇÃO ---
with st.sidebar:
    st.header("⚙️ 1. Motor de IA")
    api_key = st.text_input("Google Gemini API Key", type="password")
    
    selected_model = "models/gemini-1.5-flash" # Default fallback
    
    if api_key:
        # Seletor de Modelo
        available_models = get_available_models(api_key)
        if available_models:
            # Tenta encontrar o flash para ser o default, senão usa o primeiro da lista
            default_index = 0
            for i, m in enumerate(available_models):
                if "flash" in m:
                    default_index = i
                    break
            
            selected_model = st.selectbox(
                "Modelo Selecionado:", 
                available_models, 
                index=default_index
            )
            st.caption(f"A usar: {selected_model}")
        else:
            st.error("Chave inválida ou sem acesso a modelos.")
    else:
        st.warning("Insere a chave para iniciar.")
        st.markdown("[Obter chave gratuita](https://aistudio.google.com/)")
    
    st.divider()
    
    # --- BIBLIOTECA LEGISLATIVA DINÂMICA ---
    st.header("📚 2. Legislação Ambiental")
    st.info("Ativa os regimes legais aplicáveis:")
    
    library = legislacao.get_library()
    library_context = ""
    active_laws_count = 0
    
    for category, laws in library.items():
        with st.expander(f"📂 {category}", expanded=False):
            for law_name, details in laws.items():
                if st.checkbox(law_name, value=False, key=law_name):
                    active_laws_count += 1
                    library_context += f"- [ATIVA] {law_name} ({details['nivel']})\n"
                    library_context += f"  MANDATO: {details['mandato']}\n"
                    library_context += f"  LINK: {details['link']}\n\n"
    
    if active_laws_count > 0:
        st.success(f"✅ {active_laws_count} diplomas ativados.")

    st.divider()
    
    st.header("🌐 3. Fontes Externas")
    uploaded_legal_docs = st.file_uploader(
        "Upload PDFs Adicionais", 
        type="pdf", 
        accept_multiple_files=True,
        key=f"legal_uploader_{st.session_state.uploader_key}"
    )
    
    search_query = st.text_input("Pesquisa Web Adicional", placeholder="Ex: Portaria n.º 123/2024")
    use_web_search = st.checkbox("Ativar Pesquisa Online", value=True)
    
    st.divider()
    if st.button("🗑️ Limpar Tudo"):
        limpar_dados()

# --- FUNÇÕES ---

def get_pdf_text(pdf_file):
    text = ""
    try:
        reader = PdfReader(pdf_file)
        for page in reader.pages:
            text += page.extract_text() or ""
    except Exception as e:
        st.error(f"Erro ao ler {pdf_file.name}: {e}")
    return text

def search_online(query):
    if not query: return ""
    results_text = ""
    status = st.empty()
    status.info(f"🔎 A pesquisar: '{query}'...")
    try:
        with DDGS() as ddgs:
            results = list(ddgs.text(f"{query} legislação texto oficial", max_results=2))
        for r in results:
            try:
                page = requests.get(r['href'], timeout=4)
                soup = BeautifulSoup(page.content, 'html.parser')
                text = "\n".join([p.text for p in soup.find_all('p')])[:3000]
                results_text += f"\n>>> FONTE ONLINE: {r['title']} ({r['href']}) <<<\n{text}\n"
            except: continue
        status.empty()
        return results_text
    except Exception as e:
        status.warning(f"Erro na pesquisa web: {e}")
        return ""

def create_docx(markdown_text):
    doc = Document()
    doc.add_heading('Relatório de Análise Ambiental', 0)
    for line in markdown_text.split('\n'):
        line = line.strip()
        if not line: continue
        if line.startswith('# '): doc.add_heading(line[2:], 1)
        elif line.startswith('## '): doc.add_heading(line[3:], 2)
        elif line.startswith('### '): doc.add_heading(line[4:], 3)
        elif line.startswith('- ') or line.startswith('* '): doc.add_paragraph(line[2:], style='List Bullet')
        else: doc.add_paragraph(line.replace('**', '').replace('__', ''))
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- FUNÇÃO PRINCIPAL DE ANÁLISE ---
def run_pate_audit(target_text, lib_ctx, manual_ctx, web_ctx, api_key, model_name):
    """Executa a análise usando o modelo escolhido pelo utilizador."""
    genai.configure(api_key=api_key)
    
    # USA O MODELO SELECIONADO NA SIDEBAR
    model = genai.GenerativeModel(model_name)
    
    full_legal_context = ""
    if lib_ctx: full_legal_context += f"\n=== BIBLIOTECA LEGISLATIVA ATIVADA ===\n{lib_ctx}"
    if manual_ctx: full_legal_context += f"\n=== LEGISLAÇÃO CARREGADA MANUALMENTE ===\n{manual_ctx[:20000]}"
    if web_ctx: full_legal_context += f"\n=== PESQUISA WEB ===\n{web_ctx}"

    prompt = f"""
    Atua como um **Consultor Sénior em Ambiente, Estratégia e Ordenamento**.
    Realiza uma ANÁLISE TÉCNICA E JURÍDICA (Protocolo PATE) ao documento fornecido.

    --- BASE DE CONFORMIDADE LEGAL (A TUA "VERDADE") ---
    {full_legal_context}
    ----------------------------------------------------

    Executa estritamente o PROTOCOLO DE ANÁLISE:

    ## 1. RESUMO EXECUTIVO
    * Enquadramento do documento alvo (Objetivos, Autores, Data).
    * Maturidade do projeto/plano.

    ## 2. CHECK-UP DE CONFORMIDADE E LICENCIAMENTO
    * Cruza as medidas do documento com a 'Biblioteca Legislativa'.
    * Verifica questões críticas: **RAN/REN**, **Rede Natura 2000**, **Recursos Hídricos** e **AIA**.
    * Cita explicitamente: "A medida X alinha-se com o diploma Y" ou "Atenção: potencial conflito com Z".

    ## 3. ANÁLISE DE EXEQUIBILIDADE E DADOS
    * Avalia a qualidade da informação de base (ex: há dados de campo ou são estimativas?).
    * Avalia a viabilidade técnica das medidas propostas.

    ## 4. ANÁLISE DE RISCO
    * Riscos Ambientais (ex: impacto na biodiversidade, água).
    * Riscos de Procedimento (ex: necessidade de AIA ou AAE não identificada).

    ## 5. RECOMENDAÇÕES TÉCNICAS
    * 3 a 5 medidas corretivas concretas para garantir a aprovação e sustentabilidade do projeto.

    --- DOCUMENTO ALVO ---
    {target_text}
    """
    
    return model.generate_content(prompt).text

# --- ÁREA PRINCIPAL ---
st.subheader("📄 Documento Alvo")

uploaded_target = st.file_uploader(
    "Carrega o Relatório Técnico/EIA/Plano", 
    type="pdf",
    key=f"main_uploader_{st.session_state.uploader_key}"
)

if uploaded_target and api_key:
    if st.button("🚀 EXECUTAR ANÁLISE AMBIENTAL", type="primary"):
        with st.spinner(f"A analisar com o modelo {selected_model}..."):
            try:
                # 1. Pipeline de Texto
                target_txt = get_pdf_text(uploaded_target)
                
                manual_ctx = ""
                if uploaded_legal_docs:
                    for f in uploaded_legal_docs: manual_ctx += get_pdf_text(f)
                
                web_ctx = ""
                if use_web_search and search_query: web_ctx = search_online(search_query)
                
                # 2. Executar IA (Passando o nome do modelo)
                result = run_pate_audit(target_txt, library_context, manual_ctx, web_ctx, api_key, selected_model)
                
                # 3. Resultados
                st.success("Análise Concluída.")
                
                # Área de Limpeza
                col_clean1, col_clean2 = st.columns([1,3])
                with col_clean1:
                    if st.button("🧹 LIMPAR DADOS AGORA", type="secondary"):
                        limpar_dados()

                tab1, tab2 = st.tabs(["📝 Relatório Visual", "💾 Exportar"])
                
                with tab1: st.markdown(result)
                
                with tab2:
                    st.info("Descarrega o relatório.")
                    docx = create_docx(result)
                    st.download_button("📄 Word (.docx)", docx, f"Analise_{uploaded_target.name}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    st.download_button("📥 Markdown (.md)", result, f"Analise_{uploaded_target.name}.md")

            except Exception as e:
                st.error(f"Erro na execução: {e}")

elif not uploaded_target:
    st.info("A aguardar documento...")
